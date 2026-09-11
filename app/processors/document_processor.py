from pathlib import Path

from app.models.document_model import NormalizedDocument


def extract_document(path: Path, original_name: str) -> NormalizedDocument:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        file_type = "md" if suffix == ".markdown" else suffix[1:]
        return NormalizedDocument(file_name=original_name, file_type=file_type, text=text)
    if suffix == ".pdf":
        import fitz

        document = fitz.open(path)
        text = "\n".join(page.get_text() for page in document)
        return NormalizedDocument(file_name=original_name, file_type="pdf", text=text)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        tables = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
        return NormalizedDocument(file_name=original_name, file_type="docx", text=text, tables=tables)
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        text = "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        return NormalizedDocument(file_name=original_name, file_type="pptx", text=text)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"[{sheet.title}]")
            rows.extend(", ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
        return NormalizedDocument(file_name=original_name, file_type="xlsx", text="\n".join(rows))
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")