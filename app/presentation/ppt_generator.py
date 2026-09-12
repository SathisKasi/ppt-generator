from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from xml.etree import ElementTree

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.models.presentation_model import PresentationEnvelope
from app.presentation.layouts import SLIDE_HEIGHT, SLIDE_WIDTH
from app.presentation.slide_builder import add_bullets, add_footer, add_message, add_title
import re

from app.presentation.template_catalog import get_template
from app.presentation.themes import get_theme
from app.config import get_settings


BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)
BG_CARD = RGBColor(0x16, 0x21, 0x3E)
ACCENT_BLUE = RGBColor(0x0F, 0x34, 0x60)
ACCENT_PURPLE = RGBColor(0x53, 0x34, 0x83)
ACCENT_RED = RGBColor(0xE9, 0x45, 0x60)
TEXT_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_LIGHT = RGBColor(0xE0, 0xE0, 0xE0)
TEXT_MUTED = RGBColor(0xA0, 0xA0, 0xB0)
TEMPLATE_FONT = "Calibri"


def _pptx_theme(path: Path) -> dict:
    """Read the selected PowerPoint theme's fonts and scheme colors."""
    theme = {}
    try:
        with ZipFile(path) as archive:
            xml = ElementTree.fromstring(archive.read("ppt/theme/theme1.xml"))
        namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        colors = {}
        for color in xml.findall(".//a:clrScheme/*", namespace):
            value = next(iter(color), None)
            if value is not None:
                rgb = value.get("lastClr") or value.get("val")
                if rgb:
                    colors[color.tag.rsplit("}", 1)[-1]] = rgb
        fonts = xml.find(".//a:fontScheme/a:majorFont/a:latin", namespace)
        theme.update({
            "brand": colors.get("dk1", "301506"),
            "ink": colors.get("tx1", colors.get("dk1", "24180F")),
            "muted": colors.get("dk2", "6E6257"),
            "accent": colors.get("accent1", "FFB81C"),
            "background": colors.get("lt1", "FFFFFF"),
            "font": fonts.get("typeface") if fonts is not None else "Aptos",
        })
    except (KeyError, ElementTree.ParseError, OSError):
        return {}
    return theme


def _template_theme(template: dict) -> dict:
    configured = {key: value for key, value in template.items() if key in {"background", "ink", "muted", "accent", "brand", "font"}}
    if template["path"]:
        configured.update(_pptx_theme(template["path"]))
    return configured


def _layout_for(deck: Presentation, slide_type: str):
    if not deck.slide_layouts:
        return None
    if slide_type == "title":
        for layout in deck.slide_layouts:
            if "title" in layout.name.lower() and "content" not in layout.name.lower():
                return layout
    for layout in deck.slide_layouts:
        if "title and content" in layout.name.lower() or "title" in layout.name.lower():
            return layout
    return deck.slide_layouts[0]


def _remove_slide(deck: Presentation, slide) -> None:
    slide_ids = deck.slides._sldIdLst
    for slide_id in list(slide_ids):
        if deck.part.related_slide(slide_id.rId) is slide:
            deck.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
            return


def _clear_template_slide_text(slide) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame:
            shape.text_frame.clear()


def _template_text_shapes(slide):
    return [
        shape for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        and shape.width > Inches(1)
        and shape.height > Inches(0.15)
        and shape.top < Inches(6.8)
    ]


def _replace_shape_text(shape, text: str) -> None:
    frame = shape.text_frame
    paragraphs = list(frame.paragraphs)
    if not paragraphs:
        frame.text = text
        return
    paragraph = paragraphs[0]
    paragraph.text = text
    for extra in paragraphs[1:]:
        frame._txBody.remove(extra._p)


def _copy_template_slide(deck: Presentation, source_slide, layout) -> object:
    slide = deck.slides.add_slide(layout)
    for shape in source_slide.shapes:
        slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    return slide


def _template_prototype_index(slide_model) -> int:
    title = slide_model.title.lower()
    visual = slide_model.visual.type.lower()
    if slide_model.slide_type == "title":
        return 0
    if "comparison" in visual or "traditional" in title or "vs" in title:
        return 5
    if "timeline" in visual or "roadmap" in title:
        return 6
    if "process" in visual or "process" in title or "workflow" in title:
        return 7
    if "architecture" in visual or "architecture" in title:
        return 8
    if "metric" in visual or "performance" in title or "kpi" in title:
        return 9
    if "takeaway" in visual or "conclusion" in title or "next step" in title:
        return 10
    if "image" in visual or "diagram" in visual:
        return 4
    if "section" in visual:
        return 3
    if len(slide_model.content) >= 5:
        return 2
    return 1


def _set_textbox_font(text_frame, font_name: str, font_size: int, color_hex: str, bold: bool = False) -> None:
    for paragraph in text_frame.paragraphs:
        paragraph.font.name = font_name
        paragraph.font.size = font_size
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color_hex)


def _text_boxes(slide):
    return [
        shape for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.width > Inches(1) and shape.top < Inches(6.8)
    ]


def _fill_template_regions(slide, slide_model) -> None:
    boxes = sorted(_text_boxes(slide), key=lambda shape: (shape.top, shape.left))
    if not boxes:
        return
    title = slide_model.title
    lines = [str(item) for item in slide_model.content]
    message = slide_model.executive_message
    _replace_shape_text(boxes[0], title)

    if len(boxes) == 2:
        _replace_shape_text(boxes[1], message or "\n".join(lines))
    elif len(boxes) == 3 and boxes[-1].top > Inches(5.8):
        _replace_shape_text(boxes[1], message or "\n".join(lines))
    elif len(boxes) == 5:
        midpoint = max(1, (len(lines) + 1) // 2)
        _replace_shape_text(boxes[1], "Left")
        _replace_shape_text(boxes[2], "\n".join(lines[:midpoint]))
        _replace_shape_text(boxes[3], "Right")
        _replace_shape_text(boxes[4], "\n".join(lines[midpoint:]))
    elif len(boxes) == 6:
        midpoint = max(1, (len(lines) + 1) // 2)
        _replace_shape_text(boxes[1], "Option A")
        _replace_shape_text(boxes[2], "\n".join(lines[:midpoint]))
        _replace_shape_text(boxes[4], "Option B")
        _replace_shape_text(boxes[5], "\n".join(lines[midpoint:]))
    elif len(boxes) >= 9:
        for index, box in enumerate(boxes[1:-1], start=0):
            if index < len(lines):
                _replace_shape_text(box, lines[index])
    elif len(boxes) == 3:
        _replace_shape_text(boxes[1], message or "\n".join(lines))
    else:
        _replace_shape_text(boxes[1], "\n".join(lines) or message)


def _force_white_text(deck: Presentation) -> None:
    white = RGBColor(255, 255, 255)
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.color.rgb = white
                for run in paragraph.runs:
                    run.font.color.rgb = white


def _fill_background(slide, color: RGBColor = BG_DARK) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left: float, top: float, width: float, height: float, fill_color: RGBColor, line_color: RGBColor | None = None, line_width: float = 1.0):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def _oval(slide, left: float, top: float, width: float, height: float, fill_color: RGBColor, line_color: RGBColor | None = None):
    shape = slide.shapes.add_shape(9, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def _txt(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    color: RGBColor = TEXT_LIGHT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font_name: str = TEMPLATE_FONT,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text or "")
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _template_lines(slide_model) -> list[str]:
    lines = [str(item).strip() for item in slide_model.content if str(item).strip()]
    cleaned = []
    for line in lines:
        cleaned.append(line.lstrip("#").lstrip("*-").strip())
    if slide_model.executive_message:
        message = slide_model.executive_message.strip()
        if message and message not in cleaned:
            cleaned.insert(0, message)
    return cleaned


def _bullets(slide, items: list[str], left: float, top: float, width: float, height: float, size: int = 17, color: RGBColor | None = None, font_name: str = TEMPLATE_FONT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    bullet = chr(8226)
    text_color = color or TEXT_LIGHT
    for index, item in enumerate(items):
        text = str(item).lstrip("#").lstrip("*-").strip()
        if not text:
            continue
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(8)
        run = paragraph.add_run()
        run.text = f"{bullet} {text}"
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.color.rgb = text_color
    return box


def _header_bar(slide, title: str, title_size: int = 28) -> None:
    _rect(slide, 0, 0, 13.333, 1.4, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.08, 1.4, ACCENT_RED)
    _txt(slide, title, 0.3, 0.1, 12.5, 1.1, size=title_size, bold=True, color=TEXT_WHITE)


def _draw_reference_title(slide, slide_model) -> None:
    _fill_background(slide, BG_DARK)
    _rect(slide, 0, 6.2, 13.333, 0.1, ACCENT_RED)
    _rect(slide, 0, 0, 13.333, 0.08, ACCENT_PURPLE)
    _rect(slide, 0, 0, 0.05, 7.5, ACCENT_PURPLE)
    _txt(slide, slide_model.title, 1.0, 2.0, 11.0, 1.5, size=44, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
    subtitle = slide_model.executive_message or slide_model.purpose
    _txt(slide, subtitle, 1.0, 3.8, 11.0, 0.8, size=22, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)
    _txt(slide, "CONFIDENTIAL", 0.3, 6.8, 3.0, 0.4, size=10, color=TEXT_MUTED)


def _draw_reference_section(slide, slide_model) -> None:
    _fill_background(slide, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.08, 7.5, ACCENT_RED)
    _rect(slide, 0, 3.5, 13.333, 0.06, ACCENT_RED)
    _txt(slide, slide_model.title, 1.0, 1.8, 11.0, 1.4, size=40, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
    description = slide_model.executive_message or slide_model.purpose
    _txt(slide, description, 1.5, 3.7, 10.0, 0.8, size=18, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)


def _draw_reference_content(slide, slide_model) -> None:
    _fill_background(slide, BG_DARK)
    _header_bar(slide, slide_model.title)
    _rect(slide, 0.3, 1.6, 12.7, 5.5, BG_CARD)
    lines = _template_lines(slide_model)
    _bullets(slide, lines[:8], 0.6, 1.8, 12.2, 5.0, size=18)


def _draw_reference_two_column(slide, slide_model) -> None:
    _fill_background(slide, BG_DARK)
    _header_bar(slide, slide_model.title)
    lines = _template_lines(slide_model)
    midpoint = max(1, (len(lines) + 1) // 2)
    left_lines = lines[:midpoint]
    right_lines = lines[midpoint:] or lines[:midpoint]
    _rect(slide, 0.3, 1.6, 6.1, 5.5, BG_CARD)
    _txt(slide, "Overview", 0.4, 1.7, 5.8, 0.6, size=16, bold=True, color=ACCENT_RED)
    _bullets(slide, left_lines[:6], 0.4, 2.4, 5.8, 4.4, size=16)
    _rect(slide, 6.6, 1.6, 0.05, 5.5, ACCENT_PURPLE)
    _rect(slide, 6.9, 1.6, 6.1, 5.5, BG_CARD)
    _txt(slide, "Details", 7.0, 1.7, 5.8, 0.6, size=16, bold=True, color=ACCENT_PURPLE)
    _bullets(slide, right_lines[:6], 7.0, 2.4, 5.8, 4.4, size=16)


def _draw_reference_timeline(slide, slide_model) -> None:
    _fill_background(slide, BG_DARK)
    _header_bar(slide, slide_model.title)
    _rect(slide, 0.5, 3.9, 12.3, 0.06, ACCENT_RED)
    lines = (_template_lines(slide_model) + [""] * 4)[:4]
    positions_x = [1.0, 4.0, 7.0, 10.0]
    for index, (line, x) in enumerate(zip(lines, positions_x)):
        _oval(slide, x + 0.85, 3.65, 0.3, 0.3, ACCENT_RED)
        _txt(slide, f"Step {index + 1}", x + 0.5, 4.1, 1.3, 0.4, size=14, bold=True, color=ACCENT_RED, align=PP_ALIGN.CENTER)
        if index % 2 == 0:
            _rect(slide, x + 0.2, 2.0, 2.2, 1.5, BG_CARD)
            _txt(slide, line, x + 0.3, 2.1, 2.0, 1.2, size=13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)
        else:
            _rect(slide, x + 0.2, 4.7, 2.2, 1.5, BG_CARD)
            _txt(slide, line, x + 0.3, 4.8, 2.0, 1.2, size=13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)


def _draw_reference_conclusion(slide, slide_model) -> None:
    _fill_background(slide, BG_DARK)
    _rect(slide, 0, 0, 13.333, 2.0, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.08, 2.0, ACCENT_RED)
    _txt(slide, slide_model.title, 0.3, 0.3, 12.5, 1.4, size=30, bold=True, color=TEXT_WHITE)
    lines = _template_lines(slide_model)
    midpoint = max(1, (len(lines) + 1) // 2)
    _rect(slide, 0.3, 2.2, 7.8, 4.9, BG_CARD)
    _txt(slide, "SUMMARY", 0.4, 2.3, 3.0, 0.4, size=12, bold=True, color=ACCENT_RED)
    _bullets(slide, lines[:midpoint][:6], 0.4, 2.8, 7.5, 4.0, size=17)
    _rect(slide, 8.4, 2.2, 4.6, 4.9, ACCENT_PURPLE)
    _txt(slide, "NEXT STEPS", 8.5, 2.3, 4.3, 0.4, size=12, bold=True, color=TEXT_WHITE)
    next_steps = lines[midpoint:] or ["Validate source evidence", "Confirm decision owner", "Agree next action"]
    _txt(slide, "\n".join(f"{index + 1}. {step}" for index, step in enumerate(next_steps[:5])), 8.5, 2.8, 4.3, 3.5, size=15, color=TEXT_WHITE)
    _rect(slide, 3.0, 7.0, 7.3, 0.35, ACCENT_RED)
    _txt(slide, slide_model.executive_message or "Questions Welcome", 3.0, 7.0, 7.3, 0.35, size=13, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)


def _reference_layout_name(slide_model, index: int, total: int) -> str:
    if index == 0 or slide_model.slide_type == "title":
        return "title"
    if index == total - 1:
        return "conclusion"
    if total == 8:
        sequence = ["title", "section", "content", "section", "content", "two_column", "timeline", "conclusion"]
        return sequence[index]
    if index in {1, 3} and total >= 6:
        return "section"
    if index == total - 2 and total >= 5:
        return "timeline"
    return "content"


def _render_reference_template_deck(presentation: PresentationEnvelope, output_path: Path, template: dict) -> Path:
    deck = Presentation()
    deck.slide_width = SLIDE_WIDTH
    deck.slide_height = SLIDE_HEIGHT
    blank_layout = deck.slide_layouts[6]
    slide_models = presentation.presentation.slides
    drawers = {
        "title": _draw_reference_title,
        "section": _draw_reference_section,
        "content": _draw_reference_content,
        "two_column": _draw_reference_two_column,
        "timeline": _draw_reference_timeline,
        "conclusion": _draw_reference_conclusion,
    }
    for index, slide_model in enumerate(slide_models):
        slide = deck.slides.add_slide(blank_layout)
        layout_name = _reference_layout_name(slide_model, index, len(slide_models))
        drawers[layout_name](slide, slide_model)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.presentation.theme.update({"template_id": template["id"], "template_name": template["name"]})
    deck.save(output_path)
    return output_path


def _populate_template_slide(slide, slide_model, theme: dict) -> None:
    text_shapes = _text_boxes(slide)
    if not text_shapes:
        add_title(slide, slide_model.title, slide_model.section, theme)
        return
    _fill_template_regions(slide, slide_model)


def _populate_body_placeholder(slide, slide_model, theme: dict) -> None:
    body_placeholder = next(
        (placeholder for placeholder in slide.placeholders if placeholder.placeholder_format.type == PP_PLACEHOLDER.BODY),
        None,
    )
    if body_placeholder is None:
        return

    text_frame = body_placeholder.text_frame
    text_frame.clear()
    font_name = theme.get("font", "Aptos")
    ink_color = theme.get("ink", "1F1F1F")
    lines = []
    if slide_model.executive_message:
        lines.append(slide_model.executive_message)
    lines.extend(str(item) for item in slide_model.content)
    if not lines:
        return

    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.font.name = font_name
        paragraph.font.size = Pt(18 if len(lines) <= 3 else 15)
        paragraph.font.color.rgb = RGBColor.from_string(ink_color)
        paragraph.space_after = Pt(14)


def _render_generated_slide(slide, slide_model, theme: dict, use_template: bool) -> None:
    if use_template:
        _populate_template_slide(slide, slide_model, theme)
        return

    add_title(slide, slide_model.title, slide_model.section, theme)
    if slide_model.executive_message:
        add_message(slide, slide_model.executive_message, theme)
    if slide_model.content:
        add_bullets(slide, [str(item) for item in slide_model.content], y=Inches(2.55) if slide_model.executive_message else Inches(1.55), theme=theme)


def generate_pptx(presentation: PresentationEnvelope, output_path: Path, template_id: str = "ups-healthcare") -> Path:
    template = get_template(get_settings(), template_id)
    if template["path"]:
        return _render_reference_template_deck(presentation, output_path, template)
    theme = _template_theme(template)
    deck = Presentation()
    if not template["path"]:
        deck.slide_width = SLIDE_WIDTH
        deck.slide_height = SLIDE_HEIGHT
    slide_models = presentation.presentation.slides
    if template["path"]:
        template_slides = list(deck.slides)
        generated_slides = []
        for slide_model in slide_models:
            prototype = template_slides[min(_template_prototype_index(slide_model), len(template_slides) - 1)]
            slide = _copy_template_slide(deck, prototype, prototype.slide_layout)
            _clear_template_slide_text(slide)
            _render_generated_slide(slide, slide_model, theme, use_template=True)
            generated_slides.append(slide)
        for template_slide in template_slides:
            _remove_slide(deck, template_slide)
        slides_to_render = []
    else:
        generated_slides = []
        slides_to_render = slide_models

    for slide_model in slides_to_render:
        if template["path"]:
            slide = None
        else:
            layout = _layout_for(deck, slide_model.slide_type)
            slide = deck.slides.add_slide(layout)
            _render_generated_slide(slide, slide_model, theme, use_template=False)
        if not template["path"]:
            add_footer(slide, slide_model.slide_number, theme)
    if template["path"]:
        _force_white_text(deck)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.presentation.theme.update({"template_id": template["id"], "template_name": template["name"]})
    deck.save(output_path)
    return output_path
