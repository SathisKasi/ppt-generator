from pptx.util import Inches, Pt

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN_X = Inches(0.65)
TITLE_Y = Inches(0.45)
CONTENT_Y = Inches(1.45)


def title_style(shape) -> None:
    shape.text_frame.paragraphs[0].font.size = Pt(28)
    shape.text_frame.paragraphs[0].font.bold = True
