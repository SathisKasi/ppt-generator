from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.presentation.layouts import CONTENT_Y, MARGIN_X, TITLE_Y
from app.presentation.themes import THEME


def add_title(slide, title: str, section: str = "", theme: dict | None = None) -> None:
    theme = theme or THEME
    box = slide.shapes.add_textbox(MARGIN_X, TITLE_Y, Inches(12), Inches(0.7))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = theme["font"]
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor.from_string(theme["brand"])
    if section:
        tag = slide.shapes.add_textbox(MARGIN_X, Inches(0.16), Inches(10), Inches(0.2))
        tag.text_frame.text = section.upper()
        tag.text_frame.paragraphs[0].font.size = Pt(9)
        tag.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(THEME["muted"])


def add_footer(slide, number: int, theme: dict | None = None) -> None:
    theme = theme or THEME
    line = slide.shapes.add_shape(1, MARGIN_X, Inches(7.12), Inches(12), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string(theme["accent"])
    line.line.fill.background()
    footer = slide.shapes.add_textbox(MARGIN_X, Inches(7.18), Inches(12), Inches(0.18))
    footer.text_frame.text = f"UPS Healthcare  |  AI Presentation Generator  |  {number}"
    footer.text_frame.paragraphs[0].font.size = Pt(8)
    footer.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme["muted"])


def add_bullets(slide, items: list[str], x=MARGIN_X, y=CONTENT_Y, width=Inches(11.8), height=Inches(4.9), theme: dict | None = None) -> None:
    theme = theme or THEME
    box = slide.shapes.add_textbox(x, y, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.font.name = theme["font"]
        paragraph.font.size = Pt(18 if len(items) <= 3 else 15)
        paragraph.font.color.rgb = RGBColor.from_string(theme["ink"])
        paragraph.space_after = Pt(14)


def add_message(slide, message: str, theme: dict | None = None) -> None:
    theme = theme or THEME
    box = slide.shapes.add_textbox(MARGIN_X, CONTENT_Y, Inches(11.7), Inches(1.1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = message
    paragraph.font.name = theme["font"]
    paragraph.font.size = Pt(22)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor.from_string(theme["brand"])
    paragraph.alignment = PP_ALIGN.LEFT