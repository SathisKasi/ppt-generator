from pathlib import Path

from pptx import Presentation

from app.agents.document_agent import analyze
from app.agents.presentation_planner import plan
from app.models.document_model import NormalizedDocument
from app.presentation.ppt_generator import generate_pptx
from app.processors import extract_document


def test_document_to_pptx(tmp_path: Path):
    document = NormalizedDocument(file_name="source.txt", file_type="txt", text="Warehouse visibility\nProblem: manual updates\nObjective: improve visibility\nTarget: 20%")
    knowledge = analyze(document)
    presentation = plan(knowledge)
    output = generate_pptx(presentation, tmp_path / "result.pptx")
    assert output.exists()
    assert output.stat().st_size > 0
    assert len(presentation.presentation.slides) >= 3


def test_generated_pptx_has_one_footer_per_slide(tmp_path: Path):
    document = NormalizedDocument(file_name="source.txt", file_type="txt", text="Warehouse visibility\nProblem: manual updates\nObjective: improve visibility\nTarget: 20%")
    presentation = plan(analyze(document), 5)
    output = generate_pptx(presentation, tmp_path / "result.pptx", "clean-corporate")
    deck = Presentation(output)

    for index, slide in enumerate(deck.slides, start=1):
        footer_text = f"AI Presentation Generator  |  {index}"
        matching_shapes = [
            shape for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and footer_text in shape.text
        ]
        assert len(matching_shapes) == 1


def test_powerpoint_template_matches_reference_alignment(tmp_path: Path):
    document = NormalizedDocument(
        file_name="source.txt",
        file_type="txt",
        text="Warehouse visibility\nProblem: manual updates\nObjective: improve visibility\nTarget: 20%",
    )
    presentation = plan(analyze(document), 8)
    output = generate_pptx(presentation, tmp_path / "template-result.pptx", "presentation_template")
    expected = Presentation(Path("output/FINAL_test.pptx"))
    actual = Presentation(output)

    assert len(actual.slides) == len(expected.slides)
    for actual_slide, expected_slide in zip(actual.slides, expected.slides):
        assert len(actual_slide.shapes) == len(expected_slide.shapes)
        for actual_shape, expected_shape in zip(actual_slide.shapes, expected_slide.shapes):
            assert actual_shape.shape_type == expected_shape.shape_type
            assert actual_shape.left == expected_shape.left
            assert actual_shape.top == expected_shape.top
            assert actual_shape.width == expected_shape.width
            assert actual_shape.height == expected_shape.height


def test_supply_chain_markdown_keeps_numbered_sections():
    source_path = Path("uploads/Supply Chain Management.txt")
    knowledge = analyze(extract_document(source_path, source_path.name))
    presentation = plan(knowledge)

    assert knowledge.subject == "Supply Chain Management: Building an Efficient and Resilient Supply Network"
    assert len(knowledge.sections) == 14
    assert len(presentation.presentation.slides) == 15
    assert presentation.presentation.slides[1].title == "Introduction"
    assert presentation.presentation.slides[-1].title == "Conclusion"


def test_planner_honors_requested_slide_count():
    source_path = Path("uploads/Supply Chain Management.txt")
    knowledge = analyze(extract_document(source_path, source_path.name))

    assert len(plan(knowledge, 5).presentation.slides) == 5
    assert len(plan(knowledge, 15).presentation.slides) == 15
    assert len(plan(knowledge, 20).presentation.slides) == 20
